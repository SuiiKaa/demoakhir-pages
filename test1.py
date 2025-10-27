import os, time, signal, atexit, traceback, csv, math, json, subprocess
from pathlib import Path
from datetime import datetime
import numpy as np
import matplotlib
try:
    matplotlib.use("Qt5Agg")
except Exception:
    pass
import matplotlib.pyplot as plt

import nidaqmx
import nidaqmx.constants as C
from nidaqmx.stream_readers import AnalogSingleChannelReader
from nidaqmx.errors import DaqError

# ============= USER SETTINGS =============
DEV_CH = "cDAQ9191-1823A37Mod1/ai1"   # channel fisik accelerometer
SENS_MV_PER_G = 100.0                 # sensitivitas sensor (mV/g)
IEPE_mA = 4.0                         # current excitation mA

FS = 16800                            # sampling rate (Hz)
N = 512                               # samples dibaca per loop read_many_sample
WINDOW_SECONDS = 5.0                  # panjang window buffer untuk RMS/grafik (detik)
YRANGE_G = (-0.5, 0.5)                # range plot realtime matplotlib
BUF_SEC = 5.0                         # initial DAQ buffer (s)

# Logging lokal (Excel/CSV)
PREFER_XLSX = True                    # True = tulis XLSX
LOG_PERIOD_S = 1.0                    # tiap 1 detik log RMS
OUT_BASENAME = f"rms_log_{time.strftime('%Y%m%d_%H%M%S')}"

# >>>> Tambahan buat dashboard GitHub Pages <<<<
MAX_POINTS_JSON = 1200                # batas jumlah titik waveform yang dipublish
PUSH_INTERVAL_S = 10.0                # jeda push git (detik)
REPO_DIR = Path(r"D:\Pkl Halia\demoakhir-pages")  # lokasi repo Pages lokal
SNAPSHOT_PATH = REPO_DIR / "data" / "snapshot.json"
GIT_REMOTE = "origin"
GIT_BRANCH = "main"
# ==========================================

# Siapkan writer XLSX; kalau gagal -> fallback CSV (;)
xlsx_ok = False
wb = ws = ts_style = None
xlsx_path = os.path.join(os.getcwd(), OUT_BASENAME + ".xlsx")
csv_path  = os.path.join(os.getcwd(), OUT_BASENAME + ".csv")

if PREFER_XLSX:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import NamedStyle
        wb = Workbook()
        ws = wb.active
        ws.title = "RMS_Log"
        ws.append(["timestamp_local", "rms_g", "peak_g"])
        ts_style = NamedStyle(name="ts_ms", number_format="yyyy-mm-dd hh:mm:ss.000")
        # header bold
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
        xlsx_ok = True
    except Exception as e:
        print(f"[info] openpyxl tidak tersedia ({e}), fallback ke CSV ';'.")

RUN = True
CLOSED = False


def request_stop(*_):
    global RUN
    RUN = False


def safe_close(task):
    global CLOSED
    if CLOSED:
        return
    for action in (
        lambda: task.control(C.TaskAction.TASK_ABORT),
        task.stop,
        task.close
    ):
        try:
            action()
        except Exception:
            pass
    CLOSED = True


def wait_until_available(task, need, fs, max_wait=2.0):
    """
    Tunggu sampe buffer DAQmx punya >= need samples,
    tapi jangan bikin CPU 100%.
    """
    t0 = time.time()
    while task.in_stream.avail_samp_per_chan < need and (time.time()-t0) < max_wait:
        time.sleep(max(need/fs*0.25, 0.002))


def capture_bg(fig, ax):
    """
    Buat teknik blit (biar plot realtime smooth).
    """
    fig.canvas.draw()
    return fig.canvas.copy_from_bbox(ax.bbox)


# ====== Tambahan untuk publish snapshot.json ke GitHub Pages ======

def decimate(y: np.ndarray, fs: float, max_points: int):
    """
    Downsample array y supaya panjangnya <= max_points.
    Return y_ds, dt_ds.
    dt_ds = jarak antar titik hasil (detik).
    """
    n = len(y)
    if n <= max_points:
        step = 1
    else:
        step = math.ceil(n / max_points)  # misal 84000/1200 ~= 70
    y_ds = y[::step].copy()
    dt_ds = step / fs
    return y_ds, dt_ds


def build_snapshot_for_gweb(sig_window: np.ndarray, fs: float):
    """
    Bikin payload JSON:
    {
      "rms": <float>,
      "waveform": {
        "t0": <epoch float>,
        "dt": <detik_per_sample>,
        "y": [ ... ]
      }
    }

    - sig_window: buffer scroll_buf terakhir (WINDOW_SECONDS detik)
    - fs: sampling rate
    """
    now_epoch = time.time()
    duration_sec = len(sig_window) / fs
    t0_epoch = now_epoch - duration_sec  # kira kira waktu sample pertama window

    rms_val = float(np.sqrt(np.mean(sig_window ** 2)))

    y_ds, dt_ds = decimate(sig_window, fs, MAX_POINTS_JSON)

    snap = {
        "rms": rms_val,
        "waveform": {
            "t0": float(t0_epoch),
            "dt": float(dt_ds),
            "y": y_ds.tolist()
        }
    }
    return snap


def write_snapshot_file(snapshot: dict, path: Path):
    """
    Tulis snapshot.json ke repo GitHub Pages lokal.
    Pastikan folder data/ ada.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(snapshot, f, indent=2, sort_keys=True)


def run_git_cmd(args, cwd: Path):
    """
    Helper buat manggil git tanpa nge-crash kalau gagal.
    Balikin (rc, stdout, stderr).
    """
    proc = subprocess.run(
        ["git"] + args,
        cwd=str(cwd),
        capture_output=True,
        text=True
    )
    return proc.returncode, proc.stdout.strip(), proc.stderr.strip()


def git_publish(repo_dir: Path, snapshot_relpath: str):
    """
    Publish snapshot.json ke branch main:
    1. git pull --rebase origin main
    2. git add data/snapshot.json
    3. git commit -m "[auto] snapshot <timestamp>"
    4. git push origin main
    """
    # sync dulu biar ga divergen
    run_git_cmd(["pull", "--rebase", GIT_REMOTE, GIT_BRANCH], cwd=repo_dir)

    # stage snapshot.json
    run_git_cmd(["add", snapshot_relpath], cwd=repo_dir)

    # commit (kalau ga ada diff nyata, rc=1 "nothing to commit" -> aman)
    ts = datetime.now().isoformat(timespec="seconds")
    rc, out, err = run_git_cmd(
        ["commit", "-m", f"[auto] snapshot {ts}"],
        cwd=repo_dir
    )

    # push
    run_git_cmd(["push", GIT_REMOTE, GIT_BRANCH], cwd=repo_dir)

# ==================================================================


def main():
    global RUN
    # handle Ctrl+C, close figure, dll
    signal.signal(signal.SIGINT,  request_stop)
    signal.signal(signal.SIGTERM, request_stop)
    if hasattr(signal, "SIGBREAK"):
        signal.signal(signal.SIGBREAK, request_stop)

    task = nidaqmx.Task()
    atexit.register(lambda: safe_close(task))

    # CSV fallback kalau ga bisa XLSX
    f_csv = None
    writer = None
    if not xlsx_ok:
        f_csv = open(csv_path, "a", newline="", encoding="utf-8")
        # delimiter ';' supaya Excel Indo auto-split kolom
        writer = csv.writer(f_csv, delimiter=';')
        writer.writerow(["timestamp_local", "rms_g", "peak_g"])

    try:
        # ================== SETUP TASK NI-DAQmx ==================
        # Channel accelerometer IEPE (mode g)
        task.ai_channels.add_ai_accel_chan(
            physical_channel=DEV_CH,
            name_to_assign_to_channel="accel0",
            sensitivity=SENS_MV_PER_G,
            sensitivity_units=C.AccelSensitivityUnits.MILLIVOLTS_PER_G,
            units=C.AccelUnits.G,
            current_excit_source=C.ExcitationSource.INTERNAL,
            current_excit_val=IEPE_mA / 1000.0  # mA → A
        )
        # coupling AC kalau device support
        try:
            task.ai_channels[0].ai_coupling = C.Coupling.AC
        except Exception:
            pass

        # Timing continuous
        buf_samps = int(max(FS * BUF_SEC, N * 40))
        task.timing.cfg_samp_clk_timing(
            rate=FS,
            sample_mode=C.AcquisitionType.CONTINUOUS,
            samps_per_chan=buf_samps
        )
        try:
            task.in_stream.input_buf_size = buf_samps
        except Exception:
            pass
        try:
            task.in_stream.overwrite = C.OverwriteMode.OVERWRITE_UNREAD_SAMPLES
        except Exception:
            pass

        reader = AnalogSingleChannelReader(task.in_stream)
        block = np.zeros(N, dtype=float)

        # ================== SETUP PLOT REALTIME ==================
        plt.ion()
        fig, ax = plt.subplots(1, 1, figsize=(11, 4), constrained_layout=True)
        fig.canvas.mpl_connect("close_event", lambda evt: request_stop())

        scroll_len = int(FS * WINDOW_SECONDS)  # jumlah sample yg ditahan buat RMS/plot
        t_axis = np.arange(scroll_len) / FS
        scroll_buf = np.zeros(scroll_len, dtype=float)

        (line,) = ax.plot(t_axis, scroll_buf, animated=True)
        ax.set_xlim(t_axis[0], t_axis[-1])
        ax.set_ylim(*YRANGE_G)
        ax.set_xlabel("Time (s)")
        ax.set_ylabel("Acceleration (g)")

        text = ax.text(
            0.01, 0.95, "",
            transform=ax.transAxes,
            va="top", ha="left",
            animated=True
        )

        bg = capture_bg(fig, ax)

        def on_resize(_):
            nonlocal bg
            bg = capture_bg(fig, ax)

        fig.canvas.mpl_connect("resize_event", on_resize)

        # ================== START TASK ==================
        task.start()
        wait_until_available(task, N, FS, max_wait=2.0)

        last_log  = 0.0   # buat periodic log RMS ke XLSX / CSV
        last_save = 0.0   # periodic autosave XLSX
        last_push = 0.0   # buat periodic git push snapshot.json

        # ================== LOOP AKUISISI ==================
        while RUN and plt.fignum_exists(fig.number):
            # pastikan ada cukup sample baru
            if task.in_stream.avail_samp_per_chan < N:
                time.sleep(max(N/FS * 0.25, 0.002))
                continue

            try:
                reader.read_many_sample(
                    block,
                    number_of_samples_per_channel=N,
                    timeout=0.0
                )
            except DaqError:
                # DAQmx rewel sebentar, skip iterasi
                continue
            except Exception:
                traceback.print_exc()
                break

            # Geser buffer scrolling
            scroll_buf[:-N] = scroll_buf[N:]
            scroll_buf[-N:] = block

            # Hitung RMS & peak dari JENDELA scroll_buf (bukan cuma blok N)
            rms  = float(np.sqrt(np.mean(scroll_buf ** 2)))
            peak = float(np.max(np.abs(scroll_buf)))

            # Update tampilan realtime pakai blit
            try:
                fig.canvas.restore_region(bg)
                line.set_ydata(scroll_buf)
                text.set_text(f"RMS={rms:.4f} g   Peak={peak:.4f} g")
                ax.draw_artist(line)
                ax.draw_artist(text)
                fig.canvas.blit(ax.bbox)
                fig.canvas.flush_events()
            except Exception:
                # kalau window di-resize keras / redraw fail, ambil ulang bg
                bg = capture_bg(fig, ax)

            now = time.time()

            # ========== LOGGING PERIODIK KE XLSX / CSV ==========
            if (now - last_log) >= LOG_PERIOD_S:
                if xlsx_ok:
                    dt = datetime.now()
                    ws.append([dt, rms, peak])
                    # format kolom timestamp agar ada milidetik
                    cell = ws.cell(row=ws.max_row, column=1)
                    cell.style = ts_style
                else:
                    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                    writer.writerow([ts, rms, peak])
                    f_csv.flush()
                last_log = now

            # Autosave workbook tiap 5 detik biar aman kalau crash
            if xlsx_ok and (now - last_save) >= 5.0:
                wb.save(xlsx_path)
                last_save = now

            # ========== SNAPSHOT.JSON + GIT PUSH KE GITHUB PAGES ==========
            # 1. build payload buat dashboard G Web
            snap = build_snapshot_for_gweb(scroll_buf, FS)

            # 2. tulis ke D:\Pkl Halia\demoakhir-pages\data\snapshot.json
            write_snapshot_file(snap, SNAPSHOT_PATH)

            # 3. tiap ~10 detik, commit+push
            if (now - last_push) >= PUSH_INTERVAL_S:
                try:
                    git_publish(
                        REPO_DIR,
                        snapshot_relpath=str(SNAPSHOT_PATH.relative_to(REPO_DIR))
                    )
                except Exception as e:
                    # jangan matiin loop cuma gara2 git, cukup print
                    print("[warn] git publish error:", e)
                last_push = now

        # end while

    except Exception:
        traceback.print_exc()
    finally:
        # Tutup DAQ task rapi
        safe_close(task)

        # Matikan interactive mode matplotlib biar window bisa close normal
        try:
            plt.ioff()
            plt.show()
        except Exception:
            pass

        # Simpan file log terakhir
        if xlsx_ok:
            try:
                wb.save(xlsx_path)
                print(f"XLSX disimpan ke: {xlsx_path}")
            except Exception as e:
                print(f"[warn] gagal simpan xlsx akhir: {e}")
        else:
            try:
                if writer:
                    f_csv.flush()
                if f_csv:
                    f_csv.close()
                print(f"CSV disimpan ke: {csv_path}")
            except Exception:
                pass


if __name__ == "__main__":
    main()
